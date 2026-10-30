"""Builds the downloadable product packs into backend/files/{slug}.zip (source: user's product-pack script)."""
import os
import zipfile
import tempfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from reportlab.lib.fonts import addMapping
import arabic_reshaper
from bidi.algorithm import get_display

ROOT = Path(__file__).resolve().parent.parent
FILES = ROOT / "files"
FONTS = ROOT / "fonts"
BASE = Path(tempfile.mkdtemp(prefix="nova_products_"))

pdfmetrics.registerFont(TTFont("Nova", str(FONTS / "DejaVuSans.ttf")))
pdfmetrics.registerFont(TTFont("Nova-Bold", str(FONTS / "DejaVuSans-Bold.ttf")))
addMapping("Nova", 0, 0, "Nova")
addMapping("Nova", 1, 0, "Nova-Bold")


def ar(text):
    return get_display(arabic_reshaper.reshape(text))


black = "171717"


def style_sheet(ws, widths):
    ws.sheet_view.showGridLines = False
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor=black)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def money_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly Dashboard"
    ws.append(["NOVA MONEY PLANNER", "Amount (AED)", "Notes"])
    for r in [
        ["Monthly income", 0, "Enter total take-home income"],
        ["Fixed expenses", 0, "Rent, bills, subscriptions"],
        ["Variable expenses", 0, "Food, transport, shopping"],
        ["Debt payments", 0, "Loans and repayments"],
        ["Savings", 0, "Amount moved to savings"],
        ["Money remaining", "=B2-B3-B4-B5-B6", "Calculated automatically"],
        ["Savings rate", "=IFERROR(B6/B2,0)", "Calculated automatically"],
    ]:
        ws.append(r)
    style_sheet(ws, {"A": 28, "B": 20, "C": 42})
    for r in range(2, 8):
        ws[f"B{r}"].number_format = '"AED " #,##0.00'
        ws[f"B{r}"].fill = PatternFill("solid", fgColor="FFF2CC")
        ws[f"B{r}"].font = Font(name="Arial", bold=True, color="7F6000")
    ws["B8"].number_format = "0.0%"
    ws["B8"].fill = PatternFill("solid", fgColor="E2F0D9")
    tx = wb.create_sheet("Transactions")
    tx.append(["Date", "Type", "Category", "Description", "Amount (AED)"])
    style_sheet(tx, {"A": 16, "B": 18, "C": 22, "D": 38, "E": 18})
    for r in range(2, 102):
        tx.cell(r, 1).number_format = "yyyy-mm-dd"
        tx.cell(r, 5).number_format = '"AED " #,##0.00'
    dv = DataValidation(type="list", formula1='"Income,Expense,Savings,Debt"', allow_blank=True)
    tx.add_data_validation(dv)
    dv.add("B2:B101")
    go = wb.create_sheet("Savings Goals")
    go.append(["Goal", "Target (AED)", "Saved (AED)", "Remaining (AED)", "Target date"])
    style_sheet(go, {"A": 28, "B": 20, "C": 20, "D": 22, "E": 18})
    for r in range(2, 32):
        go.cell(r, 4, f'=IF(B{r}="","",B{r}-C{r})')
        for c in [2, 3, 4]:
            go.cell(r, c).number_format = '"AED " #,##0.00'
        go.cell(r, 5).number_format = "yyyy-mm-dd"
    wb.save(path)


def business_workbook(path):
    wb2 = Workbook()
    b = wb2.active
    b.title = "Business Snapshot"
    b.append(["NOVA BUSINESS STARTER", "Your plan"])
    for row in [
        ["Business name", ""], ["Customer problem", ""], ["Ideal customer", ""], ["Product / service", ""],
        ["Why customers choose you", ""], ["Sales channel", ""], ["Estimated price (AED)", 0],
        ["Cost per sale (AED)", 0], ["Estimated profit per sale (AED)", "=B8-B9"], ["Monthly sales target", 0],
        ["Estimated monthly gross profit (AED)", "=B10*B11"],
    ]:
        b.append(row)
    style_sheet(b, {"A": 34, "B": 62})
    for r in [8, 9, 10, 12]:
        b[f"B{r}"].number_format = '"AED " #,##0.00'
        b[f"B{r}"].fill = PatternFill("solid", fgColor="FFF2CC")
    b["B11"].number_format = "0"
    b["B11"].fill = PatternFill("solid", fgColor="FFF2CC")
    market = wb2.create_sheet("Customer Research")
    market.append(["Question", "Notes / Evidence"])
    for row in [
        "Who has this problem?", "How do they solve it today?", "What do existing alternatives cost?",
        "What do customers dislike about alternatives?", "Where can you reach 10 potential customers?",
        "What did 5 potential customers say?", "What is the smallest test offer?", "What result would validate demand?",
    ]:
        market.append([row, ""])
    style_sheet(market, {"A": 48, "B": 70})
    tasks = wb2.create_sheet("Launch Checklist")
    tasks.append(["Task", "Status", "Notes"])
    for task in [
        "Define one specific customer and problem", "Interview at least 5 potential customers",
        "Review 5 competing offers and prices", "Create a simple sample or prototype",
        "Calculate price, costs and expected margin", "Check business licensing and activity requirements",
        "Prepare terms, privacy and refund information", "Test checkout and digital delivery",
        "Make first 10 sales through legitimate channels", "Review feedback and improve the offer",
    ]:
        tasks.append([task, "Not started", ""])
    style_sheet(tasks, {"A": 60, "B": 20, "C": 42})
    status = DataValidation(type="list", formula1='"Not started,In progress,Done"', allow_blank=True)
    tasks.add_data_validation(status)
    status.add("B2:B11")
    wb2.save(path)


def habit_workbook(path):
    from openpyxl.utils import get_column_letter
    wb3 = Workbook()
    h = wb3.active
    h.title = "30 Day Tracker"
    h.append(["Habit"] + [f"Day {i}" for i in range(1, 31)] + ["Completed"])
    style_sheet(h, {"A": 28, **{get_column_letter(i): 8 for i in range(2, 32)}, "AF": 14})
    for r in range(2, 12):
        h.cell(r, 1, f"Habit {r - 1}")
        h.cell(r, 32, f'=COUNTIF(B{r}:AE{r},"✓")')
        for c in range(2, 32):
            h.cell(r, c).alignment = Alignment(horizontal="center", vertical="center")
            h.cell(r, c).fill = PatternFill("solid", fgColor="FFF8E7")
        h.cell(r, 32).alignment = Alignment(horizontal="center")
    mark = DataValidation(type="list", formula1='"✓,–"', allow_blank=True)
    h.add_data_validation(mark)
    mark.add("B2:AE11")
    notes = wb3.create_sheet("Weekly Reflection")
    notes.append(["Week", "What went well?", "What was difficult?", "One adjustment for next week"])
    for i in range(1, 5):
        notes.append([i, "", "", ""])
    style_sheet(notes, {"A": 14, "B": 42, "C": 42, "D": 42})
    wb3.save(path)


styles = getSampleStyleSheet()
title = ParagraphStyle("NTitle", parent=styles["Title"], fontName="Nova-Bold", fontSize=22, textColor=colors.HexColor("#C9A646"), alignment=TA_CENTER, spaceAfter=8)
sub = ParagraphStyle("NSub", parent=styles["Normal"], fontName="Nova", fontSize=10, textColor=colors.HexColor("#555555"), alignment=TA_CENTER, spaceAfter=18)
head = ParagraphStyle("NHead", parent=styles["Heading2"], fontName="Nova-Bold", fontSize=13, textColor=colors.HexColor("#A98225"), spaceBefore=10, spaceAfter=5)
body = ParagraphStyle("NBody", parent=styles["BodyText"], fontName="Nova", fontSize=9.5, leading=14, spaceAfter=6)
small = ParagraphStyle("NSmall", parent=body, fontSize=8, leading=11, textColor=colors.HexColor("#666666"))


def build_pdf(path, product, subtitle_en, subtitle_ar, sections, disclaimer=None):
    story = [Paragraph(product, title), Paragraph(f"{subtitle_en} | {ar(subtitle_ar)}", sub)]
    for heading, paragraphs in sections:
        story.append(Paragraph(heading, head))
        for p in paragraphs:
            story.append(Paragraph(p, body))
    if disclaimer:
        story.append(Spacer(1, 8))
        story.append(Paragraph(disclaimer, small))
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=48,
                            title=product, author="NOVA DIGITAL")
    doc.build(story)


MONEY_SECTIONS = [
    ("01 — Start here", [
        "Use the Excel workbook with Microsoft Excel or a compatible spreadsheet app. Start on the <b>Monthly Dashboard</b> tab.",
        "Enter your monthly take-home income, fixed expenses, variable expenses, debt payments and savings. The remaining balance and savings rate calculate automatically."]),
    ("02 — Track every transaction", [
        "Open <b>Transactions</b> and add one row for each payment or deposit. Select Income, Expense, Savings or Debt, then add a category, short description and amount.",
        "Use positive amounts. Keep the same currency (AED) throughout the workbook."]),
    ("03 — Build savings goals", [
        "Open <b>Savings Goals</b>. Add a goal, target amount, amount already saved and target date. The remaining amount is calculated for you."]),
    ("04 — A simple weekly routine", [
        "Once a week, enter any missing transactions, check the remaining balance, and decide one small adjustment for the next week.",
        "At month-end, compare your plan with what actually happened. Adjust next month's budget based on real spending, not guesswork."]),
    ("Important", [
        "This planner is an organizational tool, not financial, tax, investment or debt advice. Protect your file and do not share private financial information."]),
]

BUSINESS_SECTIONS = [
    ("01 — Start with a problem", [
        "Do not begin with a logo or a large inventory. Write down one clear problem experienced by one specific type of customer.",
        "Good starting point: “I help [specific customer] solve [specific problem] with [simple offer].”"]),
    ("02 — Validate before spending", [
        "Speak with at least five people who match your target customer. Ask how they solve the problem today, what is frustrating, and what they have already paid for.",
        "Look at at least five alternatives. Record their prices, strengths and customer complaints. Do not copy their work or claim their results."]),
    ("03 — Build a small offer", [
        "Create the smallest useful version of your product or service. Explain exactly what is included, who it is for, delivery time and any limits.",
        "Test interest with real conversations or a transparent waitlist. Do not advertise guaranteed outcomes."]),
    ("04 — Know your numbers", [
        "Use the Business Snapshot tab to enter your planned selling price and direct cost per sale. Estimated gross profit per sale is price minus direct cost.",
        "Monthly gross profit estimate equals gross profit per sale multiplied by your monthly sales target. This is not net profit: fees, taxes, overhead and refunds may still apply."]),
    ("05 — Launch responsibly", [
        "Before taking payments, verify the correct business activity and licence permissions for your company, and check payment-provider requirements.",
        "Prepare accurate product descriptions, contact details, delivery terms, privacy information and a clear refund policy that complies with applicable law."]),
    ("30-day action plan", [
        "Week 1: choose a customer and problem; interview five people. Week 2: compare alternatives and define your offer. Week 3: build a sample and calculate costs. Week 4: test the offer, collect feedback and improve."]),
]

HABIT_SECTIONS = [
    ("01 — Choose fewer habits", [
        "Start with one to three habits. Make each habit specific and easy to measure, such as “walk for 15 minutes” rather than “be healthier.”"]),
    ("02 — Make it achievable", [
        "Choose a small version for busy days. Example: read one page, stretch for two minutes, or prepare tomorrow's task list.",
        "Attach the habit to an existing routine, such as after breakfast or before bedtime."]),
    ("03 — Use the tracker", [
        "Open the 30 Day Tracker tab. Rename the habit rows, then mark ✓ when completed or – when skipped. The completed count updates automatically.",
        "Use Weekly Reflection to record what worked, what got in the way and one adjustment for the next week."]),
    ("04 — Review without guilt", [
        "A missed day is information, not failure. Restart at the next opportunity. Aim for steady progress rather than a perfect streak.",
        "At the end of 30 days, keep the habits that helped and simplify or replace the ones that did not."]),
]

README = """NOVA DIGITAL — {name}

Included:
- {xlsx}  (open with Microsoft Excel, Google Sheets, Numbers or a compatible app)
- {pdf}  (step-by-step guide)

Questions or a faulty file? Contact NOVA DIGITAL:
WhatsApp +971554318114 · ali94724999@icloud.com

Licence: personal use only. Please do not resell or share these files.
"""

PRODUCTS = [
    ("nova-money-planner", "NOVA MONEY PLANNER", money_workbook,
     ("A practical monthly money system", "دليل عملي لتنظيم أموالك", MONEY_SECTIONS, None)),
    ("nova-business-starter", "NOVA BUSINESS STARTER", business_workbook,
     ("A beginner-friendly business validation workbook", "دليل بداية مشروعك", BUSINESS_SECTIONS,
      "This workbook is educational and does not guarantee sales or business success. Check applicable UAE licensing, consumer-protection, tax and data-protection requirements before trading.")),
    ("nova-habit-tracker", "NOVA HABIT TRACKER", habit_workbook,
     ("30-day consistency system", "نظام متابعة لمدة 30 يوماً", HABIT_SECTIONS, None)),
]


def main():
    FILES.mkdir(exist_ok=True)
    for slug, name, make_xlsx, (sub_en, sub_ar, sections, disclaimer) in PRODUCTS:
        stem = name.replace(" ", "_")
        xlsx, pdf = BASE / f"{stem}.xlsx", BASE / f"{stem}_GUIDE.pdf"
        make_xlsx(xlsx)
        build_pdf(pdf, name, sub_en, sub_ar, sections, disclaimer)
        out = FILES / f"{slug}.zip"
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(xlsx, xlsx.name)
            z.write(pdf, pdf.name)
            z.writestr("START_HERE.txt", README.format(name=name, xlsx=xlsx.name, pdf=pdf.name))
        stale = FILES / f"{slug}.pdf"
        if stale.exists():
            stale.unlink()
        print(out.name, os.path.getsize(out), "bytes")


if __name__ == "__main__":
    main()
